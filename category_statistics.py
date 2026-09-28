"""Export report-ready product counts by source category."""

from __future__ import annotations

import datetime as dt
import json
import os
from collections import defaultdict
from pathlib import Path
import sqlite3
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


PLATFORM_LABELS = {"bdt": "八达通", "sjyx": "中建三局严选", "lcgt": "乐从钢铁"}
PLATFORM_ORDER = ("bdt", "sjyx", "lcgt")
HEADERS = ["平台", "一级分类", "二级分类", "三级分类", "最小分类商品数量"]


def _lcgt_categories(path: Path) -> dict[str, tuple[str, str]]:
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    categories: dict[str, tuple[str, str]] = {}
    if not isinstance(data, dict):
        return categories
    for level1_name, level2_items in data.items():
        if not isinstance(level2_items, dict):
            continue
        for level2_name, item in level2_items.items():
            if isinstance(item, dict) and item.get("subcategory_id") is not None:
                categories[str(item["subcategory_id"])] = (str(level1_name), str(level2_name))
    return categories


def _sjyx_rows(connection: sqlite3.Connection) -> list[tuple[str, str, str, str, int]]:
    return [
        ("sjyx", level1 or "未分类", level2 or "", level3 or "", count)
        for level1, level2, level3, count in connection.execute(
            """
            SELECT json_extract(payload_json, '$.category.names[0]'),
                   json_extract(payload_json, '$.category.names[1]'),
                   json_extract(payload_json, '$.category.names[2]'),
                   COUNT(*)
            FROM products
            WHERE source = 'sjyx' AND status = 'success'
            GROUP BY 1, 2, 3
            """
        )
    ]


def _bdt_rows(connection: sqlite3.Connection) -> list[tuple[str, str, str, str, int]]:
    table_exists = connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'bdt_category_mappings'"
    ).fetchone()
    if not table_exists:
        return [
            ("bdt", "未分类", "", "", count)
            for (count,) in connection.execute(
                "SELECT COUNT(*) FROM products WHERE source = 'bdt' AND status = 'success'"
            )
        ]
    return [
        ("bdt", level1 or "未分类", level2 or "", level3 or "", count)
        for level1, level2, level3, count in connection.execute(
            """
            SELECT mapping.level1_name, mapping.level2_name, mapping.level3_name, COUNT(*)
            FROM products AS product
            LEFT JOIN bdt_category_mappings AS mapping
              ON product.category_id = mapping.category_id
            WHERE product.source = 'bdt' AND product.status = 'success'
            GROUP BY 1, 2, 3
            """
        )
    ]


def _lcgt_rows(connection: sqlite3.Connection, mapping_path: Path) -> list[tuple[str, str, str, str, int]]:
    categories = _lcgt_categories(mapping_path)
    grouped: dict[tuple[str, str], int] = defaultdict(int)
    for category_id, count in connection.execute(
        """
        SELECT category_id, COUNT(*)
        FROM products
        WHERE source = 'lcgt' AND status = 'success'
        GROUP BY category_id
        """
    ):
        level1, level2 = categories.get(str(category_id or ""), ("未分类", ""))
        grouped[(level1, level2)] += count
    return [("lcgt", level1, level2, "", count) for (level1, level2), count in grouped.items()]


def category_rows(root: Path | str) -> list[tuple[str, str, str, str, int]]:
    root = Path(root)
    database = root / "data" / "catalog.sqlite3"
    if not database.exists():
        raise ValueError(f"找不到商品数据库：{database}")
    with sqlite3.connect(database) as connection:
        rows = _sjyx_rows(connection) + _lcgt_rows(connection, root / "lcgt" / "lcgt_products.json") + _bdt_rows(connection)
    return sorted(rows, key=lambda item: (PLATFORM_ORDER.index(item[0]), item[1], item[2], item[3]))


def export_category_statistics(root: Path | str) -> dict[str, Any]:
    root = Path(root)
    rows = category_rows(root)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    output_dir = root / "data" / "cloud-products" / f"{stamp}-category-statistics"
    output_dir.mkdir(parents=True, exist_ok=False)
    destination = output_dir / "分类数据统计.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "分类数据统计"
    sheet.append(HEADERS)
    for source, level1, level2, level3, count in rows:
        sheet.append([PLATFORM_LABELS[source], level1, level2, level3, count])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
        row[4].alignment = Alignment(horizontal="right", vertical="center")
        row[4].number_format = "#,##0"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{max(1, len(rows) + 1)}"
    for column, width in {"A": 18, "B": 18, "C": 20, "D": 22, "E": 20}.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 24

    temporary = destination.with_suffix(".tmp")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, destination)
    return {"path": str(destination), "rows": len(rows)}
