"""Build cloud-product integration workbooks from the crawler catalog."""

from __future__ import annotations

from copy import copy
import datetime as dt
from html import unescape
from html.parser import HTMLParser
import json
import os
from pathlib import Path
import re
import shutil
import threading
import time
import uuid
from typing import Any, Callable, Dict, List, Optional, Sequence

from openpyxl import load_workbook

from catalog_store import CatalogStore


PLATFORM_LABELS = {"bdt": "八达通", "sjyx": "中建三局严选", "lcgt": "乐从钢铁"}
PLATFORM_ORDER = ("bdt", "sjyx", "lcgt")
PROCESSING_MODES = {"full", "incremental"}
MODE_LABELS = {"full": "全量", "incremental": "增量"}
ROWS_PER_WORKBOOK = 1000
PROCESSING_RULE_VERSION = "2026.09.24-mode-files-v5"
ACTIVE_PROCESSING_STATUSES = {"preparing", "running", "finalizing"}
ILLEGAL_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


class ImageSourceParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.sources: List[str] = []

    def handle_starttag(self, tag: str, attrs: List[tuple[str, Optional[str]]]) -> None:
        if tag.lower() != "img":
            return
        source = dict(attrs).get("src")
        if source:
            self.sources.append(source)


def clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = ILLEGAL_EXCEL_CHARACTERS.sub("", str(value)).strip()
    return text[:32767] if text else None


def clean_parameter_part(value: Any) -> str:
    return unescape(str(value or "")).replace("：", ":").replace("；", ";").strip()


def parameter_pair(name: Any, value: Any) -> Optional[str]:
    key = clean_parameter_part(name)
    item = clean_parameter_part(value)
    return f"{key}:{item}" if key and item else None


def parameter_fragments(name: Any, value: Any) -> List[str]:
    """Build semicolon-delimited parameters and discard fragments without key:value."""
    pair = parameter_pair(name, value)
    if not pair:
        return []
    valid = []
    for fragment in pair.split(";"):
        key, separator, item = fragment.partition(":")
        key, item = key.strip(), item.strip()
        if separator and key and item:
            valid.append(f"{key}:{item}")
    return valid


def absolute_sjyx_image(value: Any) -> Optional[str]:
    source = clean_text(value)
    if source and source.startswith("/"):
        return "https://oss.3jyx.cn" + source
    return source


def detail_image_sources(html: Any) -> List[str]:
    parser = ImageSourceParser()
    parser.feed(str(html or ""))
    return [source for source in (clean_text(item) for item in parser.sources) if source]


def dictionary(value: Any) -> Dict[str, Any]:
    """Return nested API data as a mapping, tolerating null/malformed entries."""
    return value if isinstance(value, dict) else {}


def first_bdt_sku(payload: Dict[str, Any]) -> tuple[Dict[str, Any], Dict[str, Any]]:
    sku = next((item for item in payload.get("spec_value_list") or [] if isinstance(item, dict)), {})
    units = sku.get("unit_list") or []
    valid_units = [item for item in units if isinstance(item, dict)]
    unit = next((item for item in valid_units if item.get("unit_level") == 1), valid_units[0] if valid_units else {})
    return sku, unit


def map_bdt(payload: Dict[str, Any]) -> List[Any]:
    payload = dictionary(payload)
    sku, unit = first_bdt_sku(payload)
    spec = clean_text(sku.get("spec_value_str"))
    spec_definition = next((item for item in payload.get("spec_value") or [] if isinstance(item, dict)), {})
    spec_name = spec_definition.get("name") or "型号"
    images = [item for item in (clean_text(value) for value in payload.get("goods_image") or []) if item]
    return [
        None,
        clean_text(payload.get("name")),
        None,
        "优质",
        None,
        "纵购",
        spec,
        clean_text(unit.get("name")),
        None,
        "否",
        clean_text(";".join(parameter_fragments(spec_name, spec))),
        clean_text(payload.get("image")),
        clean_text(";".join(images)),
    ]


def map_sjyx(payload: Dict[str, Any]) -> List[Any]:
    payload = dictionary(payload)
    product = dictionary(payload.get("product"))
    detail = dictionary(product.get("detail"))
    brand = clean_text(product.get("brand"))
    if not brand or brand == "无品牌":
        brand = "优质"
    parameters = []
    for attribute in product.get("attributes") or []:
        if not isinstance(attribute, dict):
            continue
        parameters.extend(parameter_fragments(attribute.get("attributeName"), attribute.get("attributeValue")))
    images = detail_image_sources(detail.get("detailContentOfHtml"))
    return [
        None,
        clean_text(product.get("name")),
        None,
        brand,
        None,
        "纵购",
        clean_text(product.get("spec_model")),
        clean_text(product.get("unit")),
        None,
        "否",
        clean_text(";".join(parameters)),
        absolute_sjyx_image(detail.get("mainImage")),
        clean_text(";".join(images)),
    ]


def map_lcgt(payload: Dict[str, Any]) -> List[Any]:
    payload = dictionary(payload)
    parameters = []
    for name, value in (("材质", payload.get("material")), ("钢厂", payload.get("steel_mill"))):
        parameters.extend(parameter_fragments(name, value))
    return [
        None,
        clean_text(payload.get("name")),
        None,
        "优质",
        None,
        "纵购",
        clean_text(payload.get("spec")),
        clean_text(payload.get("unit")),
        None,
        "否",
        clean_text(";".join(parameters)),
        clean_text(payload.get("main_image") or payload.get("image")),
        None,
    ]


MAPPERS: Dict[str, Callable[[Dict[str, Any]], List[Any]]] = {
    "bdt": map_bdt,
    "sjyx": map_sjyx,
    "lcgt": map_lcgt,
}


def write_workbook(template: Path, destination: Path, rows: Sequence[Sequence[Any]]) -> None:
    """Copy the user template, remove its rule rows, and write one data part."""
    workbook = load_workbook(template)
    sheet = workbook.active
    template_styles = [copy(sheet.cell(2, column)._style) for column in range(1, 14)]
    template_alignments = [copy(sheet.cell(2, column).alignment) for column in range(1, 14)]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell._style = copy(template_styles[column_index - 1])
            alignment = copy(template_alignments[column_index - 1])
            alignment.wrap_text = column_index not in (12, 13)
            alignment.vertical = "top"
            cell.alignment = alignment
        sheet.row_dimensions[row_index].height = 42
    last_row = max(1, len(rows) + 1)
    for table in sheet.tables.values():
        table.ref = f"A1:M{last_row}"
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, destination)


class CloudProductProcessor:
    def __init__(self, root: Path | str) -> None:
        self.root = Path(root)
        self.database = self.root / "data" / "catalog.sqlite3"
        self.template = self.root / "data" / "samples" / "云商品批量导入模板.xlsx"
        self.output_root = self.root / "data" / "cloud-products"
        self.lock = threading.RLock()
        self.jobs: Dict[str, Dict[str, Any]] = {}
        self._restore()

    def _restore(self) -> None:
        metadata_root = self.output_root / ".jobs"
        if not metadata_root.exists():
            return
        for path in metadata_root.glob("*.json"):
            try:
                job = json.loads(path.read_text(encoding="utf-8"))
                if job.get("status") in ACTIVE_PROCESSING_STATUSES:
                    job.update(status="failed", ended=time.time(), error="服务重启，原数据处理任务已中断")
                self.jobs[job["id"]] = job
            except (OSError, ValueError, KeyError):
                continue

    def _persist(self, job: Dict[str, Any]) -> None:
        folder = self.output_root / ".jobs"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{job['id']}.json"
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(job, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def snapshot(self) -> Dict[str, Any]:
        with self.lock:
            jobs = sorted(self.jobs.values(), key=lambda item: item["started"], reverse=True)
            return {"jobs": [dict(job) for job in jobs[:20]]}

    def active_targets(self) -> set[str]:
        with self.lock:
            return {
                target
                for job in self.jobs.values()
                if job["status"] in ACTIVE_PROCESSING_STATUSES
                for target in job["targets"]
            }

    def start(self, targets: Any, mode: Any) -> str:
        if not isinstance(targets, list) or not targets:
            raise ValueError("请选择至少一个数据平台")
        if any(not isinstance(target, str) or target not in PLATFORM_LABELS for target in targets):
            raise ValueError("请选择有效的数据平台")
        if mode not in PROCESSING_MODES:
            raise ValueError("请选择全量处理或增量处理")
        targets = [target for target in PLATFORM_ORDER if target in set(targets)]
        with self.lock:
            if any(job["status"] in ACTIVE_PROCESSING_STATUSES for job in self.jobs.values()):
                raise ValueError("已有云商品数据处理任务正在运行")
            stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
            job_id = f"{stamp}-{mode}-{uuid.uuid4().hex[:6]}"
            job = {
                "id": job_id,
                "mode": mode,
                "targets": targets,
                "status": "preparing",
                "started": time.time(),
                "ended": None,
                "total": 0,
                "processed": 0,
                "percent": 0,
                "generated_files": 0,
                "rule_version": PROCESSING_RULE_VERSION,
                "current_source": None,
                "output_dir": None,
                "error": None,
            }
            self.jobs[job_id] = job
            self._persist(job)
            threading.Thread(target=self._run, args=(job_id,), daemon=True).start()
            return job_id

    def _update(self, job: Dict[str, Any], *, persist: bool = False, **values: Any) -> None:
        with self.lock:
            job.update(values)
            total = int(job.get("total") or 0)
            processed = int(job.get("processed") or 0)
            job["percent"] = 100 if job.get("status") == "success" else (round(processed * 100 / total, 1) if total else 0)
            if persist:
                self._persist(job)

    def _run(self, job_id: str) -> None:
        job = self.jobs[job_id]
        mode = job["mode"]
        targets = job["targets"]
        output_name = job_id
        temporary_dir = self.output_root / ("." + output_name + ".tmp")
        final_dir = self.output_root / output_name
        try:
            if not self.template.exists():
                raise ValueError(f"找不到云商品模板：{self.template}")
            if not self.database.exists():
                raise ValueError(f"找不到商品数据库：{self.database}")
            with CatalogStore(self.database) as store:
                totals = {source: store.processing_count(source, mode) for source in targets}
                self._update(job, total=sum(totals.values()), status="running", persist=True)
                temporary_dir.mkdir(parents=True, exist_ok=False)
                job_template = temporary_dir / ".template.xlsx"
                shutil.copy2(self.template, job_template)
                summary = {
                    "id": job_id,
                    "mode": mode,
                    "rule_version": PROCESSING_RULE_VERSION,
                    "platforms": {},
                    "total": job["total"],
                }
                processed_since_update = 0
                for source in targets:
                    self._update(job, current_source=source, persist=True)
                    rows: List[List[Any]] = []
                    part = 1
                    source_count = 0
                    source_files: List[str] = []
                    for payload in store.iter_processing_payloads(source, mode):
                        rows.append(MAPPERS[source](payload))
                        source_count += 1
                        processed_since_update += 1
                        if len(rows) == ROWS_PER_WORKBOOK:
                            filename = self._write_part(job, source, part, rows, temporary_dir, job_template)
                            source_files.append(filename)
                            rows = []
                            part += 1
                        if processed_since_update >= 50:
                            processed_since_update = 0
                            self._update(job, processed=job["processed"] + 50)
                    if processed_since_update:
                        self._update(job, processed=job["processed"] + processed_since_update)
                        processed_since_update = 0
                    if rows:
                        filename = self._write_part(job, source, part, rows, temporary_dir, job_template)
                        source_files.append(filename)
                    summary["platforms"][source] = {
                        "label": PLATFORM_LABELS[source],
                        "products": source_count,
                        "files": source_files,
                    }
                self._update(job, status="finalizing", current_source=None, persist=True)
                summary["generated_files"] = job["generated_files"]
                (temporary_dir / "processing-summary.json").write_text(
                    json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                manifest = [
                    f"处理模式：{MODE_LABELS[mode]}处理",
                    f"处理批次：{job_id}",
                    f"规则版本：{PROCESSING_RULE_VERSION}",
                    f"商品总数：{job['total']}",
                    f"表格数量：{job['generated_files']}",
                    "",
                    "待导入文件：",
                ]
                for source in targets:
                    manifest.extend(summary["platforms"][source]["files"])
                (temporary_dir / "待导入文件清单.txt").write_text(
                    "\n".join(manifest) + "\n", encoding="utf-8"
                )
                job_template.unlink()
                self.output_root.mkdir(parents=True, exist_ok=True)
                temporary_dir.replace(final_dir)
                store.mark_processing_completed(targets, job_id)
            self._update(
                job,
                status="success",
                ended=time.time(),
                processed=job["total"],
                output_dir=str(final_dir),
                current_source=None,
                persist=True,
            )
        except Exception as exc:  # noqa: BLE001 - surface background failures in the UI.
            if temporary_dir.exists():
                shutil.rmtree(temporary_dir, ignore_errors=True)
            self._update(
                job,
                status="failed",
                ended=time.time(),
                error=str(exc),
                current_source=None,
                persist=True,
            )

    def _write_part(
        self,
        job: Dict[str, Any],
        source: str,
        part: int,
        rows: Sequence[Sequence[Any]],
        folder: Path,
        template: Path,
    ) -> str:
        platform_folder = folder / PLATFORM_LABELS[source]
        mode_label = MODE_LABELS[job["mode"]]
        filename = f"{PLATFORM_LABELS[source]}_{mode_label}_{part:03d}.xlsx"
        write_workbook(template, platform_folder / filename, rows)
        self._update(job, generated_files=job["generated_files"] + 1)
        return str(Path(PLATFORM_LABELS[source]) / filename)
