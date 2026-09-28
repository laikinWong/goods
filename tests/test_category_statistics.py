import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import load_workbook

from catalog_store import CatalogStore
from category_statistics import export_category_statistics


class CategoryStatisticsTests(unittest.TestCase):
    def test_export_groups_each_platform_at_its_available_category_depth(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            database = root / "data" / "catalog.sqlite3"
            with CatalogStore(database) as store:
                run_id = store.start_run("sjyx")
                store.mark_success("sjyx", "s1", {"category": {"names": ["土建材料", "钢筋类", "螺纹钢"]}}, run_id)
                store.mark_success("sjyx", "s2", {"category": {"names": ["土建材料", "钢筋类", "螺纹钢"]}}, run_id)
                store.finish_run(run_id)
                run_id = store.start_run("bdt")
                store.mark_success("bdt", "b1", {"name": "铁钉"}, run_id, "1001")
                store.finish_run(run_id)
                run_id = store.start_run("lcgt")
                store.mark_success("lcgt", "l1", {"name": "工字钢"}, run_id, "100")
                store.finish_run(run_id)
            with sqlite3.connect(database) as connection:
                connection.execute(
                    "INSERT INTO bdt_category_mappings VALUES "
                    "('1001','857','钉线建材','1000','手动钉类','铁钉','2026-09-28T00:00:00+00:00')"
                )
            mapping = {"型材": {"工字钢": {"subcategory_id": "100"}}}
            mapping_path = root / "lcgt" / "lcgt_products.json"
            mapping_path.parent.mkdir()
            mapping_path.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")

            result = export_category_statistics(root)
            workbook = load_workbook(result["path"], read_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
            workbook.close()

            self.assertEqual(rows[0], ("平台", "一级分类", "二级分类", "三级分类", "最小分类商品数量"))
            self.assertIn(("中建三局严选", "土建材料", "钢筋类", "螺纹钢", 2), rows)
            self.assertIn(("八达通", "钉线建材", "手动钉类", "铁钉", 1), rows)
            self.assertIn(("乐从钢铁", "型材", "工字钢", None, 1), rows)


if __name__ == "__main__":
    unittest.main()
