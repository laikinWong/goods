import json
from pathlib import Path
import tempfile
import time
import unittest

from openpyxl import Workbook, load_workbook

from catalog_store import CatalogStore
from cloud_product_processing import CloudProductProcessor, map_bdt, map_lcgt, map_sjyx


HEADERS = [
    "商品编码（选填，留空自动生成）", "商品名称*", "物料号", "品牌", "分类编码",
    "来源（纵购/B2）", "单位规格", "基本单位", "基本单位价格", "是否自营（是/否）",
    "参数（格式：颜色：白色；长度：10m）", "主图（内嵌图片或图片URL）", "详情图（内嵌图片或图片URL）",
]


class MappingTests(unittest.TestCase):
    def test_platform_rules_use_values_and_ascii_parameter_separators(self):
        bdt = {
            "name": "胶带", "image": "main.jpg", "goods_image": ["1.jpg", "2.jpg"],
            "spec_value": [{"name": "型号"}],
            "spec_value_list": [{"spec_value_str": "40×15", "unit_list": [{"unit_level": 1, "name": "卷"}]}],
        }
        sjyx = {"product": {"name": "钢筋", "brand": "无品牌", "spec_model": "M1", "unit": "根",
                 "attributes": [{"attributeName": "产地", "attributeValue": "国产"}],
                 "detail": {"mainImage": "/a.png", "detailContentOfHtml": '<img src="b.png"><img src="c.png">'}}}
        lcgt = {"name": "耐候钢", "spec": "5*C", "unit": "吨", "material": "Q355", "steel_mill": "宝钢", "main_image": "m.jpg"}

        self.assertEqual(map_bdt(bdt)[3:13], ["优质", None, "纵购", "40×15", "卷", None, "否", "型号:40×15", "main.jpg", "1.jpg;2.jpg"])
        self.assertEqual(map_sjyx(sjyx)[3], "优质")
        self.assertEqual(map_sjyx(sjyx)[10:], ["产地:国产", "https://oss.3jyx.cn/a.png", "b.png;c.png"])
        self.assertEqual(map_lcgt(lcgt)[10:], ["材质:Q355;钢厂:宝钢", "m.jpg", None])

    def test_mappers_tolerate_null_nested_api_items(self):
        bdt = {"spec_value": [None], "spec_value_list": [None, {"unit_list": [None]}]}
        sjyx = {"product": {"attributes": [None], "detail": None}}

        self.assertEqual(map_bdt(bdt)[10], None)
        self.assertEqual(map_sjyx(sjyx)[10:], [None, None, None])
        self.assertEqual(map_lcgt(None)[1], None)

    def test_sjyx_parameters_only_keep_key_value_fragments(self):
        payload = {
            "product": {
                "attributes": [
                    {"attributeName": "颜色", "attributeValue": "黄黑(可按项目指定颜色)"},
                    {"attributeName": "备注", "attributeValue": "包含螺栓及配件；包含运输；不含安装"},
                    {"attributeName": None, "attributeValue": "无参数名"},
                    {"attributeName": "空值", "attributeValue": None},
                    {"attributeName": "高*宽", "attributeValue": "1.2*2m"},
                ]
            }
        }

        self.assertEqual(
            map_sjyx(payload)[10],
            "颜色:黄黑(可按项目指定颜色);备注:包含螺栓及配件;高*宽:1.2*2m",
        )

    def test_all_platform_parameters_only_keep_key_value_fragments(self):
        bdt = {
            "spec_value": [{"name": "包装"}],
            "spec_value_list": [{"spec_value_str": "14&quot;手板锯；无参数名"}],
        }
        lcgt = {"material": "Q355；附加说明", "steel_mill": "宝钢"}

        self.assertEqual(map_bdt(bdt)[10], '包装:14"手板锯')
        self.assertEqual(map_lcgt(lcgt)[10], "材质:Q355;钢厂:宝钢")


class ProcessingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        samples = self.root / "data" / "samples"
        samples.mkdir(parents=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        sheet.append(["八达通规则"] * 13)
        sheet.append(["三局严选规则"] * 13)
        sheet.append(["乐从钢铁规则"] * 13)
        workbook.save(samples / "云商品批量导入模板.xlsx")
        workbook.close()

    def tearDown(self):
        self.tmp.cleanup()

    def wait(self, processor, job_id):
        deadline = time.monotonic() + 20
        while time.monotonic() < deadline:
            job = processor.jobs[job_id]
            if job["status"] not in {"preparing", "running", "finalizing"}:
                return job
            time.sleep(.02)
        self.fail("processing job did not finish")

    def test_incremental_processing_splits_at_1000_and_marks_crawl_rows(self):
        database = self.root / "data" / "catalog.sqlite3"
        with CatalogStore(database) as store:
            run_id = store.start_run("bdt")
            for index in range(1001):
                store.mark_success("bdt", index, {"id": index, "name": f"商品{index}"}, run_id)
            store.finish_run(run_id)
            self.assertEqual(store.processing_count("bdt", "incremental"), 1001)

        processor = CloudProductProcessor(self.root)
        job_id = processor.start(["bdt"], "incremental")
        job = self.wait(processor, job_id)

        self.assertEqual(job["status"], "success", job.get("error"))
        self.assertEqual(job["processed"], 1001)
        self.assertEqual(job["generated_files"], 2)
        self.assertEqual(job["rule_version"], "2026.09.24-mode-files-v5")
        output = Path(job["output_dir"])
        prefix = "八达通_增量"
        first = load_workbook(output / "八达通" / f"{prefix}_001.xlsx", read_only=True)
        second = load_workbook(output / "八达通" / f"{prefix}_002.xlsx", read_only=True)
        self.assertEqual(first.active.max_row, 1001)
        self.assertEqual(second.active.max_row, 2)
        self.assertEqual(first.active["B2"].value, "商品0")
        first.close()
        second.close()
        summary = json.loads((output / "processing-summary.json").read_text(encoding="utf-8"))
        self.assertEqual(summary["rule_version"], "2026.09.24-mode-files-v5")
        self.assertEqual(summary["platforms"]["bdt"]["products"], 1001)
        manifest = (output / "待导入文件清单.txt").read_text(encoding="utf-8")
        self.assertIn("处理模式：增量处理", manifest)
        self.assertIn(f"八达通/{prefix}_001.xlsx", manifest)
        self.assertIn(f"八达通/{prefix}_002.xlsx", manifest)
        with CatalogStore(database) as store:
            self.assertEqual(store.processing_count("bdt", "incremental"), 0)


if __name__ == "__main__":
    unittest.main()
